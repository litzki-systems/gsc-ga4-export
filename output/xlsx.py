"""
XLSX output helpers: write_sheet, save_workbook_tabs, save_csv.
© Litzki Systems LLC
"""
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import REPORT_BRAND

HEADER_FILL = PatternFill("solid", fgColor="1e3a5f")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial")
ALT_FILL    = PatternFill("solid", fgColor="EEF3FF")
WARN_FILL   = PatternFill("solid", fgColor="FFF3CD")
BAD_FILL    = PatternFill("solid", fgColor="FFE0E0")
GOOD_FILL   = PatternFill("solid", fgColor="D4EDDA")

NUM_FMT = {
    "ctr":              "0.00%",
    "ctr_pct":          "0.00%",
    "gsc_ctr":          "0.00%",
    "position":         "0.0",
    "gsc_position":     "0.0",
    "prev_position":    "0.0",
    "delta_position":   "0.0",
    "clicks":           "#,##0",
    "impressions":      "#,##0",
    "gsc_clicks":       "#,##0",
    "gsc_impressions":  "#,##0",
    "prev_clicks":      "#,##0",
    "prev_impressions": "#,##0",
    "sessions":         "#,##0",
    "pageviews":        "#,##0",
    "perf_score":       "0",
    "lcp_ms":           "#,##0",
    "fcp_ms":           "#,##0",
    "inp_ms":           "#,##0",
    "avg_duration_s":   "0.0",
    "bounce_rate":      "0.0",
    "engagement_rate":  "0.0",
}

MOM_DELTA_COLS = {"delta_clicks", "delta_impressions"}

def write_sheet(ws, rows, headers, highlight_fn=None):
    ws.freeze_panes = "A2"
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    for i, row in enumerate(rows, start=2):
        ws.append([row.get(h, "") for h in headers])
        fill = highlight_fn(row) if highlight_fn else (ALT_FILL if i % 2 == 0 else None)
        for cell, h in zip(ws[i], headers):
            if fill:
                cell.fill = fill
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
            if h in NUM_FMT:
                cell.number_format = NUM_FMT[h]
            if h in MOM_DELTA_COLS and isinstance(cell.value, (int, float)):
                if cell.value > 0:
                    cell.font = Font(name="Arial", size=10, color="155724", bold=True)
                elif cell.value < 0:
                    cell.font = Font(name="Arial", size=10, color="A00000", bold=True)

    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            (len(str(ws.cell(r, col_idx).value or "")) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


def write_summary(ws, tabs):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    ws["A1"] = REPORT_BRAND
    ws["A1"].font = Font(name="Arial", bold=True, size=13, color="1e3a5f")
    ws["A1"].fill = PatternFill("solid", fgColor="EEF3FF")
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 24

    ws["A2"] = "Sheet"
    ws["B2"] = "Rows"
    ws["C2"] = "Clicks"
    ws["D2"] = "Impressions"
    for cell in ws[2]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for i, entry in enumerate(tabs, start=3):
        name, rows = entry[0], entry[1]
        clicks = sum(r.get("clicks", 0) or 0 for r in rows)
        impr   = sum(r.get("impressions", 0) or 0 for r in rows)
        ws.cell(i, 1, name).font = Font(name="Arial", size=10)
        ws.cell(i, 2, len(rows)).font = Font(name="Arial", size=10)
        ws.cell(i, 2).alignment = Alignment(horizontal="center")
        ws.cell(i, 3, clicks if clicks else "").number_format = "#,##0"
        ws.cell(i, 3).font = Font(name="Arial", size=10)
        ws.cell(i, 4, impr if impr else "").number_format = "#,##0"
        ws.cell(i, 4).font = Font(name="Arial", size=10)
        if i % 2 == 0:
            for c in ws[i]:
                c.fill = ALT_FILL

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16


def save_workbook_tabs(tabs, path, log):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_summary = wb.create_sheet(title="00 Summary")
    write_summary(ws_summary, tabs)

    for entry in tabs:
        name, rows, headers = entry[0], entry[1], entry[2]
        hfn = entry[3] if len(entry) > 3 else None
        ws  = wb.create_sheet(title=name[:31])
        write_sheet(ws, rows, headers, hfn)

    wb.save(path)
    log(f"    Saved: {path.name} ({len(tabs)+1} sheets)")
    return path


def save_csv(rows, headers, path, log):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
    log(f"    Saved: {path.name} ({len(rows)} rows)")
    return path
